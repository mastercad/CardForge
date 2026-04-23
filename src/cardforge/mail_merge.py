"""
Mail-Merge-Dialog: CSV/Excel einlesen, Platzhalter ersetzen.
"""

from __future__ import annotations

import copy

from PySide6.QtWidgets import (
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)

from .models import ELEMENT_TEXT, CardLayout


def _load_csv(path: str) -> tuple[list[str], list[dict]]:
    import csv

    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = list(reader.fieldnames or [])
        for row in reader:
            rows.append(dict(row))
    return headers, rows


def _load_excel(path: str) -> tuple[list[str], list[dict]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows_raw = list(ws.iter_rows(values_only=True))
    if not rows_raw:
        return [], []
    headers = [str(c) if c is not None else "" for c in rows_raw[0]]
    rows = []
    for row in rows_raw[1:]:
        rows.append(
            {h: (str(v) if v is not None else "") for h, v in zip(headers, row, strict=False)}
        )
    return headers, rows


def _apply_merge(template_layout: CardLayout, row: dict) -> CardLayout:
    """Ersetzt {{Platzhalter}} in Text-Elementen mit Werten aus der Zeile."""
    new_layout = copy.deepcopy(template_layout)
    for elems in (new_layout.front_elements, new_layout.back_elements):
        for e in elems:
            if e.type == ELEMENT_TEXT:
                for key, val in row.items():
                    e.text = e.text.replace(f"{{{{{key}}}}}", val)
    return new_layout


class MailMergeDialog(QDialog):
    def __init__(self, template_layout: CardLayout, parent=None):
        super().__init__(parent)
        self.setWindowTitle(self.tr("Mail Merge"))
        self.setMinimumSize(620, 500)
        self._template = template_layout
        self._headers: list[str] = []
        self._rows: list[dict] = []
        self._result_layouts: list[CardLayout] = []
        self._build_ui()

    def _build_ui(self):
        root = QVBoxLayout(self)

        # Load data source
        grp_file = QGroupBox(self.tr("Data Source"))
        hl = QHBoxLayout(grp_file)
        self._path_lbl = QLabel(self.tr("No file loaded"))
        btn_load = QPushButton(self.tr("Open CSV/Excel…"))
        btn_load.clicked.connect(self._load_file)
        hl.addWidget(self._path_lbl, 1)
        hl.addWidget(btn_load)
        root.addWidget(grp_file)

        # Data preview
        grp_preview = QGroupBox(self.tr("Data Preview"))
        vl = QVBoxLayout(grp_preview)
        self._table = QTableWidget()
        vl.addWidget(self._table)
        root.addWidget(grp_preview, 1)

        # Available placeholders
        grp_ph = QGroupBox(self.tr("Available Placeholders (use in text elements)"))
        vl2 = QVBoxLayout(grp_ph)
        self._ph_list = QListWidget()
        vl2.addWidget(self._ph_list)
        root.addWidget(grp_ph)

        # Buttons
        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btns.button(QDialogButtonBox.StandardButton.Ok).setText(self.tr("Generate Cards"))
        btns.accepted.connect(self._generate)
        btns.rejected.connect(self.reject)
        root.addWidget(btns)

    def _load_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, self.tr("Open File"), "", self.tr("Spreadsheets (*.csv *.xlsx *.xls)")
        )
        if not path:
            return
        try:
            if path.lower().endswith(".csv"):
                self._headers, self._rows = _load_csv(path)
            else:
                self._headers, self._rows = _load_excel(path)
            self._path_lbl.setText(path)
            self._populate_table()
            self._populate_placeholders()
        except Exception as e:
            QMessageBox.critical(self, self.tr("Error"), str(e))

    def _populate_table(self):
        self._table.setColumnCount(len(self._headers))
        self._table.setHorizontalHeaderLabels(self._headers)
        self._table.setRowCount(min(len(self._rows), 50))
        for r, row in enumerate(self._rows[:50]):
            for c, h in enumerate(self._headers):
                self._table.setItem(r, c, QTableWidgetItem(row.get(h, "")))

    def _populate_placeholders(self):
        self._ph_list.clear()
        for h in self._headers:
            self._ph_list.addItem(QListWidgetItem(f"{{{{{h}}}}}"))

    def _generate(self):
        if not self._rows:
            QMessageBox.warning(self, self.tr("No Data"), self.tr("Please load a file first."))
            return
        self._result_layouts = []
        for i, row in enumerate(self._rows):
            layout = _apply_merge(self._template, row)
            layout.name = f"{self._template.name} [{i + 1}]"
            self._result_layouts.append(layout)
        self.accept()

    def result_layouts(self) -> list[CardLayout]:
        return self._result_layouts
